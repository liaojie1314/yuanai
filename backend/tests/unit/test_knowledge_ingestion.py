"""结构感知解析、质量检查与 OCR 降级的单元覆盖。"""

from __future__ import annotations

import importlib
import io
import logging

import pytest
from docx import Document
from openpyxl import Workbook

from app.services.knowledge_ingestion import check_quality
from app.services.knowledge_ocr import (
    OcrUnavailableError,
    load_ocr_backend,
    resolve_ocr_backend,
)
from app.services.knowledge_parse import (
    MAX_PARSED_CHARS,
    UnsupportedDocumentError,
    parse_document,
)


def test_markdown_blocks_keep_full_heading_path() -> None:
    """Markdown 按标题切块，section 保留层级路径而不是只留末级标题。"""

    content = (
        "# 测试文档1\n\n开篇说明。\n\n"
        "## 第一节\n\n第一节正文。\n\n"
        "### 细节\n\n细节正文。\n\n"
        "## 第二节\n\n第二节正文。\n"
    )
    parsed = parse_document(content.encode(), "text/markdown", "测试文档1.md")
    assert parsed.parser == "markdown"
    assert [block.section for block in parsed.blocks] == [
        "测试文档1",
        "测试文档1 > 第一节",
        "测试文档1 > 第一节 > 细节",
        "测试文档1 > 第二节",
    ]
    assert parsed.blocks[2].content == "细节正文。"


def test_html_headings_become_sections_and_scripts_are_dropped() -> None:
    """HTML 按标题切块，脚本与样式内容不进入正文。"""

    html = (
        "<html><head><style>.a{color:red}</style></head><body>"
        "<h1>报告标题</h1><p>第一段。</p>"
        "<h2>子章节</h2><p>第二段。</p>"
        "<script>window.x=1</script>"
        "</body></html>"
    )
    parsed = parse_document(html.encode(), "text/html", "report.html")
    assert parsed.parser == "html"
    assert [block.section for block in parsed.blocks] == ["报告标题", "报告标题 > 子章节"]
    joined = "\n".join(block.content for block in parsed.blocks)
    assert "window.x" not in joined
    assert "color:red" not in joined


def test_docx_heading_styles_drive_sections() -> None:
    """docx 的标题样式还原为切块边界。"""

    document = Document()
    document.add_heading("测试文档2", level=1)
    document.add_paragraph("正文第一段。")
    document.add_heading("附录", level=2)
    document.add_paragraph("附录正文。")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_document(buffer.getvalue(), "", "测试文档2.docx")
    assert parsed.parser == "docx"
    assert [block.section for block in parsed.blocks] == ["测试文档2", "测试文档2 > 附录"]


def test_spreadsheet_splits_by_sheet_and_row_range() -> None:
    """表格按工作表与行区域切块，不同工作表不混进同一个片段。"""

    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "数据表A"
    for index in range(3):
        first.append([f"行{index}", index])
    second = workbook.create_sheet("数据表B")
    second.append(["合计", 3])
    buffer = io.BytesIO()
    workbook.save(buffer)

    parsed = parse_document(buffer.getvalue(), "", "报表.xlsx")
    assert parsed.parser == "xlsx"
    assert [block.section for block in parsed.blocks] == [
        "工作表 数据表A 第 1-3 行",
        "工作表 数据表B 第 1-1 行",
    ]
    assert "合计" not in parsed.blocks[0].content


def test_python_source_splits_by_top_level_symbol() -> None:
    """Python 源码按 AST 顶层符号切块，模块级代码单独成块。"""

    source = (
        "import os\n\n"
        "CONSTANT = 1\n\n"
        "def first() -> int:\n    return 1\n\n"
        "class Second:\n    value = 2\n"
    )
    parsed = parse_document(source.encode(), "text/x-python", "sample.py")
    assert parsed.parser == "python"
    assert [block.section for block in parsed.blocks] == [
        "sample.py",
        "sample.py::first",
        "sample.py::Second",
    ]
    assert "import os" in parsed.blocks[0].content
    assert parsed.blocks[1].content.startswith("def first()")


def test_unsupported_mime_type_is_rejected() -> None:
    """未知二进制类型不入库，由调用方记成 UNSUPPORTED_FORMAT。"""

    with pytest.raises(UnsupportedDocumentError):
        parse_document(b"\x00\x01", "application/x-binary", "blob.bin")


def test_parsed_content_is_capped() -> None:
    """超长文档按上限截断，不把超出部分写进片段。"""

    parsed = parse_document(("长" * (MAX_PARSED_CHARS + 500)).encode(), "text/plain", "big.txt")
    assert sum(len(block.content) for block in parsed.blocks) == MAX_PARSED_CHARS


def test_image_uses_injected_ocr_backend() -> None:
    """图片走 OCR 路径并标记 ocr_used，测试注入假后端不加载真实模型。"""

    parsed = parse_document(b"fake-image", "image/png", "扫描件.png", ocr=lambda _: "识别出的文字")
    assert parsed.parser == "image"
    assert parsed.ocr_used is True
    assert parsed.ocr_missing is False
    assert parsed.blocks[0].content == "识别出的文字"


def test_image_without_ocr_backend_degrades_instead_of_raising() -> None:
    """OCR 后端缺失时图片解析不抛栈，返回空内容并标记 ocr_missing。"""

    parsed = parse_document(b"fake-image", "image/jpeg", "扫描件.jpg", ocr=None)
    assert parsed.parser == "image"
    assert parsed.blocks == []
    assert parsed.ocr_missing is True
    assert parsed.ocr_used is False


def test_resolve_ocr_backend_logs_warning_when_rapidocr_missing(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """未安装 RapidOCR 时只记 warning 并返回 None，不抛栈、不静默。"""

    def missing(name: str) -> object:
        raise ModuleNotFoundError(f"No module named {name!r}")

    monkeypatch.setattr(importlib, "import_module", missing)
    with pytest.raises(OcrUnavailableError):
        load_ocr_backend()
    with caplog.at_level(logging.WARNING, logger="app.services.knowledge_ocr"):
        assert resolve_ocr_backend() is None
    assert "OCR 不可用" in caplog.text


def test_quality_check_rejects_empty_and_garbled_content() -> None:
    """空内容与整体乱码判为致命问题，轻微乱码和过短内容只告警。"""

    assert check_quality([(None, "   ")]).fatal == "EMPTY_CONTENT"
    assert check_quality([(None, "�" * 10 + "ok")]).fatal == "GARBLED_CONTENT"
    short = check_quality([(None, "短文本")])
    assert short.fatal is None
    assert short.warnings == ["SHORT_CONTENT"]
    partial = check_quality([(None, "正常内容" * 20 + "�")])
    assert partial.fatal is None
    assert partial.warnings == ["PARTIALLY_GARBLED"]
