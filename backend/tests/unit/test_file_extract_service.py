"""文件上下文提取单元测试。"""

import io

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.services.file_extract_service import extract_preview, preview_context


def test_extracts_plain_text_with_bounded_context() -> None:
    preview = extract_preview(b"hello yuanai", "text/plain", "note.txt")
    assert preview.kind == "text"
    assert preview_context(preview) == "hello yuanai"


def test_extracts_csv_as_table_context() -> None:
    preview = extract_preview(b"name,score\nAda,100\n", "text/csv", "scores.csv")
    assert preview.kind == "table"
    assert preview.rows == [["name", "score"], ["Ada", "100"]]
    assert "Ada\t100" in preview_context(preview)


def test_extracts_docx_text() -> None:
    document = Document()
    document.add_paragraph("document context")
    payload = io.BytesIO()
    document.save(payload)
    preview = extract_preview(payload.getvalue(), "application/octet-stream", "note.docx")
    assert preview.kind == "text"
    assert preview.text == "document context"


def test_extracts_xlsx_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["item", "count"])
    sheet.append(["files", 2])
    payload = io.BytesIO()
    workbook.save(payload)
    preview = extract_preview(payload.getvalue(), "application/octet-stream", "files.xlsx")
    assert preview.kind == "table"
    assert preview.rows == [["item", "count"], ["files", "2"]]


def test_marks_pdf_for_visual_preview_and_text_context() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    payload = io.BytesIO()
    writer.write(payload)

    preview = extract_preview(payload.getvalue(), "application/pdf", "requirements.pdf")

    assert preview.kind == "pdf"
    assert preview_context(preview) == ""


def test_reports_unsupported_binary_as_not_previewable() -> None:
    preview = extract_preview(b"\x00\x01", "application/zip", "archive.zip")
    assert preview.kind == "unsupported"
    assert preview.supported is False
    assert preview_context(preview) == ""


def test_invalid_document_payload_degrades_to_unsupported_preview() -> None:
    preview = extract_preview(b"not a real PDF", "application/pdf", "broken.pdf")

    assert preview.kind == "unsupported"
    assert preview.supported is False
